#!/usr/bin/env python3
"""Convert Snap DSA XLSX reports to CSV files for DuckDB ingestion."""

import os
from pathlib import Path
from openpyxl import load_workbook

DATA_DIR = Path(__file__).parent / "data"

def xlsx_to_csvs(xlsx_path, output_dir):
    """Convert each sheet of an XLSX to a separate CSV."""
    wb = load_workbook(xlsx_path, read_only=True)
    os.makedirs(output_dir, exist_ok=True)

    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        csv_name = f"{sheet_name}.csv"
        csv_path = output_dir / csv_name
        rows_written = 0

        with open(csv_path, "w") as f:
            for row in ws.iter_rows():
                cells = []
                for cell in row:
                    val = cell.value
                    if val is None:
                        cells.append("")
                    else:
                        s = str(val).replace('"', '""')
                        if "," in s or '"' in s or "\n" in s:
                            s = f'"{s}"'
                        cells.append(s)
                # Skip fully empty rows
                if all(c == "" for c in cells):
                    continue
                f.write(",".join(cells) + "\n")
                rows_written += 1

        print(f"  {sheet_name}: {rows_written} rows -> {csv_path.name}")

    wb.close()

def main():
    xlsx_files = list((DATA_DIR / "eu_dsa").glob("*.xlsx"))
    if not xlsx_files:
        print("No XLSX files found in data/eu_dsa/")
        return

    for xlsx_path in sorted(xlsx_files):
        stem = xlsx_path.stem
        out_dir = DATA_DIR / "eu_dsa_csv" / stem
        print(f"\nConverting {xlsx_path.name}:")
        xlsx_to_csvs(xlsx_path, out_dir)

if __name__ == "__main__":
    main()
